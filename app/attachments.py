from __future__ import annotations

import base64
import hashlib
import io
import json
import shutil
import stat
import subprocess
import sys
import threading
import zipfile
from pathlib import Path

from PIL import Image, ImageOps

from .config import ROOT, read_json, settings, write_json
from .store import uid
from .runtime import worker_command

Image.MAX_IMAGE_PIXELS = 24_000_000
ALLOWED = {".txt", ".md", ".docx", ".doc", ".xlsx", ".xls", ".pptx", ".ppt", ".jpg", ".jpeg", ".png", ".webp"}
IMAGE_TYPES = {".jpg", ".jpeg", ".png", ".webp"}
OCR_LOCK = threading.Lock()
ENGINE = None
ARCHIVE_ITEMS = 10000
ARCHIVE_BYTES = 100 * 1024 * 1024
ARCHIVE_MEMBER_BYTES = 20 * 1024 * 1024


def archive_entries(archive):
    """Validate a ZIP directory without extracting or executing its entries."""
    entries = archive.infolist()
    if len(entries) > ARCHIVE_ITEMS or sum(v.file_size for v in entries) > ARCHIVE_BYTES:
        raise ValueError("ZIP 条目超过 1 万个或总展开量超过 100 MiB，请拆分后上传")
    seen = set()
    for item in entries:
        # ZipInfo normalizes Windows separators and truncates NULs in filename;
        # inspect the original header value before those transformations.
        name = item.orig_filename
        parts = name.rstrip("/").split("/")
        if (not name or name.startswith("/") or "\\" in name or ":" in name or
                any(part in ("", ".", "..") for part in parts) or "\x00" in name):
            raise ValueError("ZIP 含无效或越界条目路径")
        key = name.rstrip("/").casefold()
        if key in seen:
            raise ValueError("ZIP 含重复条目，无法唯一追溯来源")
        seen.add(key)
        if item.flag_bits & 1 or stat.S_ISLNK(item.external_attr >> 16):
            raise ValueError("不接受密码保护或符号链接 ZIP 条目")
        if item.file_size > ARCHIVE_MEMBER_BYTES:
            raise ValueError("ZIP 单条目展开超过 20 MiB，请拆分资料")
        if item.file_size > 1024 * 1024 and item.file_size > max(1, item.compress_size) * 200:
            raise ValueError("ZIP 压缩比异常（超过 200 倍），未展开")
    return entries


def raw_source(ident):
    meta = load(ident)
    folder = directory(ident).resolve()
    name = meta.get("source", "")
    if not name or name != Path(name).name or ":" in name or "\\" in name:
        raise ValueError("附件原件路径无效")
    source = folder / name
    if source.is_symlink() or (hasattr(source, "is_junction") and source.is_junction()) or not source.is_file():
        raise ValueError("附件原件缺失或是链接")
    if not source.resolve().is_relative_to(folder):
        raise ValueError("附件原件越界")
    if source.stat().st_size > settings()["upload_mb"] * 1024 * 1024:
        raise ValueError("附件原件超过当前上传限制")
    return meta, source


def decode_text(blob):
    """Prefer explicit BOM/strict UTF; avoid silently decoding binaries as Latin text."""
    if blob.startswith((b"PK\x03\x04", b"MZ", b"\xd0\xcf\x11\xe0", b"%PDF", b"\x89PNG", b"\xff\xd8\xff")):
        raise ValueError("内容不像纯文本，文件扩展名可能不正确；请按实际格式上传")
    if blob.startswith((b"\xff\xfe\x00\x00", b"\x00\x00\xfe\xff")):
        candidates = ["utf-32"]
    elif blob.startswith((b"\xff\xfe", b"\xfe\xff")):
        candidates = ["utf-16"]
    elif blob.startswith(b"\xef\xbb\xbf"):
        candidates = ["utf-8-sig"]
    elif b"\x00" in blob[:4096]:
        sample = blob[:4096]
        even = sample[::2].count(0) / max(1, len(sample[::2]))
        odd = sample[1::2].count(0) / max(1, len(sample[1::2]))
        if max(even, odd) < .35 or min(even, odd) > .1:
            raise ValueError("文件含无规律 NUL 字节，疑似二进制；未强行按 UTF-16 读取")
        candidates = ["utf-16-le" if odd > even else "utf-16-be"]
    else:
        candidates = ["utf-8", "gb18030"]
    for encoding in candidates:
        try:
            text = blob.decode(encoding)
            controls = sum(ord(c) < 32 and c not in "\n\r\t\f" for c in text)
            if controls > max(1, len(text) * .01) or "\x00" in text:
                continue
            return text.replace("\r\n", "\n").replace("\r", "\n"), encoding
        except UnicodeError:
            pass
    raise ValueError("无法可靠识别文本编码，或内容含大量二进制控制字符。请另存为 UTF-8/UTF-16 后上传；未用乱码替换原文。")


def directory(ident):
    if len(ident) != 32 or any(c not in "0123456789abcdef" for c in ident):
        raise ValueError("附件 ID 无效")
    return ROOT / "data/attachments" / ident


def load(ident):
    value = read_json(directory(ident) / "metadata.json")
    if not value:
        raise KeyError("附件不存在")
    return value


def save(meta):
    write_json(directory(meta["id"]) / "metadata.json", meta)
    return meta


def ocr_image(image):
    global ENGINE
    import numpy as np
    from rapidocr_onnxruntime import RapidOCR
    with OCR_LOCK:
        if ENGINE is None:
            import rapidocr_onnxruntime
            packaged = Path(rapidocr_onnxruntime.__file__).parent
            destination = ROOT / "runtime/ocr"
            for model in packaged.rglob("*.onnx"):
                dest = destination / model.name
                if not dest.exists():
                    shutil.copy2(model, dest)
            models = list(destination.glob("*.onnx"))
            kwargs = {"intra_op_num_threads": 2, "inter_op_num_threads": 1}
            for name, keyword in (("det_model_path", "det"), ("rec_model_path", "rec"), ("cls_model_path", "cls")):
                match = next((p for p in models if keyword in p.name), None)
                if match:
                    kwargs[name] = str(match)
            ENGINE = RapidOCR(**kwargs)
        result, _ = ENGINE(np.array(image.convert("RGB")))
    lines = result or []
    return "\n".join(str(row[1]) for row in lines), (sum(float(row[2]) for row in lines) / len(lines) if lines else None)


def prepare_image(blob, dest):
    with Image.open(io.BytesIO(blob)) as opened:
        if opened.width * opened.height > Image.MAX_IMAGE_PIXELS:
            raise ValueError("图片超过 2400 万像素，请先缩小")
        image = ImageOps.exif_transpose(opened).convert("RGB")
    image.thumbnail((2400, 2400))
    image.save(dest, "JPEG", quality=90)
    return image


def extract(ident, original_path=None):
    meta = load(ident)
    folder = directory(ident)
    source = folder / meta["source"]
    try:
        suffix = source.suffix.lower()
        confidence = None
        warning = ""
        office_types = {".docx", ".doc", ".xlsx", ".xls", ".pptx", ".ppt"}
        if suffix == ".zip":
            with zipfile.ZipFile(source) as archive:
                entries = archive_entries(archive)
                files = [v for v in entries if not v.is_dir()]
            text = "ZIP 原始资料包；未将条目正文注入聊天。使用附件读取工具列条目/读文字，或 Python 的 INPUTS 读取原件副本。\n"
            text += "\n".join(v.filename + "（" + str(v.file_size) + " 字节）" for v in files[:100])
            if len(files) > 100:
                text += "\n仅展示前 100 项，完整条目清单可由附件读取工具分页查看。"
            meta.update(raw_archive=True, archive_files=len(files))
            method = "ZIP 原始附件 / 安全清单（未解压）"
        elif suffix in {".docx", ".xlsx", ".pptx"} and meta.get("read_mode") != "office":
            with zipfile.ZipFile(source) as archive:
                if len(archive.infolist()) > 10000 or sum(a.file_size for a in archive.infolist()) > 100 * 1024 * 1024:
                    raise ValueError("Office 文件展开后过大，请选择较小的文件")
        if suffix == ".zip":
            pass
        elif suffix in office_types and meta.get("read_mode") == "office":
            proc = subprocess.run(worker_command("doc", original_path or source), cwd=ROOT,
                                  capture_output=True, timeout=45, creationflags=getattr(subprocess, "CREATE_NO_WINDOW", 0))
            if proc.returncode:
                raise ValueError("Office COM 读取失败。需本机安装对应 Office 且有文件读取权限；透明加密文件须能在 Office 中正常打开。" + proc.stderr.decode("utf-8", errors="replace")[-160:])
            text, method = proc.stdout.decode("utf-8"), "Office COM 只读兼容提取（经授权应用接口）"
        elif suffix in IMAGE_TYPES:
            image = prepare_image(source.read_bytes(), folder / "vision.jpg")
            text, confidence = ocr_image(image)
            method = "RapidOCR / ONNX · 中文、英文"
            meta["image"] = True
            if not text.strip():
                warning = "未识别到文字，仍可发送原图"
        elif suffix in (".txt", ".md") or suffix not in ALLOWED:
            text, encoding = decode_text(source.read_bytes())
            method = "本地文本解析 / " + encoding
            if encoding == "gb18030":
                warning = "编码推测为 GB18030，请预览核对；若乱码可另存 UTF-8 后上传"
        elif suffix == ".docx":
            from docx import Document
            from docx.oxml.ns import qn
            with zipfile.ZipFile(source) as archive:
                if sum(a.file_size for a in archive.infolist()) > 100 * 1024 * 1024:
                    raise ValueError("Word 展开后过大")
            doc = Document(source)
            parts = []
            for node in doc.element.body:
                if node.tag == qn("w:p"):
                    parts.append("".join(t.text or "" for t in node.iter(qn("w:t"))))
                elif node.tag == qn("w:tbl"):
                    for row in node.findall(qn("w:tr")):
                        parts.append(" | ".join("".join(t.text or "" for t in cell.iter(qn("w:t"))) for cell in row.findall(qn("w:tc"))))
            with zipfile.ZipFile(source) as archive:
                images = [i for i in archive.infolist() if i.filename.startswith("word/media/") and Path(i.filename).suffix.lower() in IMAGE_TYPES]
                for number, item in enumerate(images[:10], 1):
                    try:
                        image = prepare_image(archive.read(item), folder / f"embedded-{number}.jpg")
                        result, _ = ocr_image(image)
                        parts.append(f"[内嵌图片 {number} OCR]\n{result or '未识别到文字'}")
                    except Exception:
                        parts.append(f"[内嵌图片 {number} 无法识别]")
                if len(images) > 10:
                    warning = "仅 OCR 前 10 张内嵌图片"
            text, method = "\n\n".join(parts), "DOCX 正文/表格 + 内嵌图片 OCR"
        elif suffix == ".xlsx":
            import openpyxl
            book = openpyxl.load_workbook(source, read_only=True, data_only=True, keep_links=False)
            parts = []
            cells = 0
            for sheet in book:
                parts.append(f"[工作表：{sheet.title}]")
                for row in sheet.iter_rows():
                    if any(c.value is not None for c in row):
                        parts.append(" | ".join(str(c.value) if c.value is not None else "" for c in row))
                    cells += len(row)
                    if cells > 100000:
                        warning = "已达到 10 万单元格提取上限，请选择较小文件"
                        break
                if cells > 100000:
                    break
            book.close()
            text, method = "\n".join(parts), "Python / openpyxl 工作表文字（公式使用已缓存结果）"
        elif suffix == ".xls":
            import xlrd
            book = xlrd.open_workbook(str(source), on_demand=True)
            parts = []
            for sheet in book.sheets():
                parts.append(f"[工作表：{sheet.name}]")
                for row in range(min(sheet.nrows, 10000)):
                    parts.append(" | ".join(str(x) for x in sheet.row_values(row)))
                if sheet.nrows > 10000:
                    warning = "每个工作表最多提取 1 万行"
            book.release_resources()
            text, method = "\n".join(parts), "Python / xlrd 工作表文字"
        elif suffix == ".pptx":
            from pptx import Presentation
            presentation = Presentation(source)
            parts = []
            def shape_text(shapes):
                for shape in shapes:
                    if shape.shape_type == 6:
                        yield from shape_text(shape.shapes)
                    if shape.has_text_frame:
                        yield shape.text
                    if shape.has_table:
                        for row in shape.table.rows:
                            yield " | ".join(cell.text for cell in row.cells)
            for number, slide in enumerate(presentation.slides, 1):
                parts.append(f"[幻灯片 {number}]\n" + "\n".join(shape_text(slide.shapes)))
            text, method = "\n\n".join(parts), "Python / python-pptx 幻灯片文字和表格"
        else:
            raise ValueError("旧版 DOC/PPT 默认 Python 解析暂不支持。请在附件读取模式中选择 Office COM 兼容后重新上传。")
        (folder / "extracted.txt").write_text(text, encoding="utf-8")
        meta.update(status="ready", text=text, method=method, confidence=confidence, warning=warning,
                    chars=len(text), version=1)
        return save(meta)
    except Exception as error:
        meta.update(status="error", error=str(error)[:400], text="")
        return save(meta)


def upload(filename, blob, read_mode="native"):
    name = filename.replace("\\", "/").rsplit("/", 1)[-1]
    suffix = Path(name).suffix.lower()
    if len(blob) > settings()["upload_mb"] * 1024 * 1024:
        raise ValueError("文件超过设置的大小限制")
    if suffix == ".zip":
        with zipfile.ZipFile(io.BytesIO(blob)) as archive:
            archive_entries(archive)
    elif suffix not in ALLOWED:
        decode_text(blob)  # Unknown extensions are data, never executed or blindly decoded.
    if len(suffix) > 20 or any(c not in ".abcdefghijklmnopqrstuvwxyz0123456789_-" for c in suffix):
        suffix = ".txt"
    if read_mode != "office" and suffix in (".docx", ".xlsx", ".pptx") and not blob.startswith(b"PK"):
        raise ValueError("文档内容不是有效的现代 Office 压缩格式；若为企业加密文件，请切换 Office COM 兼容模式")
    ident = uid()
    folder = directory(ident)
    folder.mkdir()
    (folder / ("source" + suffix)).write_bytes(blob)
    meta = {"id": ident, "name": name[:180], "source": "source" + suffix,
            "sha256": hashlib.sha256(blob).hexdigest(), "size": len(blob),
            "status": "processing", "image": suffix in IMAGE_TYPES, "version": 0, "read_mode": read_mode}
    return save(meta)


def snapshot(ident, text=None, mode="both"):
    meta = load(ident)
    if meta["status"] != "ready" and not (meta.get("image") and (directory(ident) / "vision.jpg").exists()):
        raise ValueError(f'附件尚未准备好：{meta["name"]}')
    if mode not in ("both", "text", "image"):
        raise ValueError("无效附件模式")
    return {"id": ident, "name": meta["name"], "text": meta.get("text", "") if text is None else text,
            "mode": mode, "image": meta["image"], "method": meta.get("method", ""),
            "sha256": meta["sha256"], "version": meta["version"], "raw_archive": bool(meta.get("raw_archive"))}


def content(message, vision=False):
    parts = [message["content"]]
    images = []
    for attachment in message.get("attachments", []):
        if attachment["mode"] != "image":
            text = attachment.get("injected_text", attachment["text"])
            marker = "｜已截取部分文字，不能声称已阅读全文" if attachment.get("injection", {}).get("truncated") else ""
            parts.append(f'\n[附件：{attachment["name"]}｜编号 {attachment["id"]}｜{attachment["method"]}{marker}]\n{text}\n[/附件]')
        if vision and attachment["image"] and attachment["mode"] != "text":
            images.append(base64.b64encode((directory(attachment["id"]) / "vision.jpg").read_bytes()).decode())
    return "\n".join(parts), images


def read_authorized(args, scope=None):
    """Read only IDs assembled by the server from the active conversation branch."""
    scope = scope or {}
    allowed = list(dict.fromkeys(scope.get("attachment_ids", [])))
    ident, action = args.get("attachment_id"), args.get("action", "list")
    if not ident:
        if action != "list":
            raise ValueError("读取需要 attachment_id；先用 {\"action\":\"list\"} 查看本对话附件编号")
        items = []
        for aid in allowed:
            try:
                meta = load(aid)
                items.append({"id": aid, "name": meta["name"], "bytes": meta["size"],
                              "kind": "zip" if meta.get("raw_archive") else "image" if meta.get("image") else "text",
                              "sha256": meta["sha256"]})
            except (KeyError, OSError, ValueError):
                continue
        return {"items": items, "count": len(items), "complete": True,
                "next_step": "用 attachment_id 读取；ZIP 先列条目再指定 member。未列出的附件未获本对话授权。"}
    if ident not in allowed:
        raise ValueError("此附件不在当前对话已发送的授权附件中；请先上传并发送，不接受其他对话或任意路径")
    meta, source = raw_source(ident)
    with source.open("rb") as reader:
        if hashlib.file_digest(reader, "sha256").hexdigest() != meta["sha256"]:
            raise ValueError("附件原件校验变化，请重新上传；未使用变化后的内容")
    offset, limit = args.get("offset", 0), args.get("max_chars", 4000)
    result = {"attachment_id": ident, "name": meta["name"], "sha256": meta["sha256"]}
    member = args.get("member")
    if source.suffix.lower() == ".zip":
        with zipfile.ZipFile(source) as archive:
            entries = [v for v in archive_entries(archive) if not v.is_dir()]
            if action == "list":
                return {**result, "items": [{"member": v.filename, "bytes": v.file_size} for v in entries[offset:offset+100]],
                        "total": len(entries), "offset": offset,
                        "next_offset": offset + 100 if offset + 100 < len(entries) else None,
                        "complete": offset == 0 and len(entries) <= 100,
                        "next_step": "read 时传 attachment_id 和清单内精确 member；offset 改为文字字符位置。"}
            entry = next((v for v in entries if v.filename == member), None)
            if entry is None:
                raise ValueError("请提供 ZIP 清单内确切的 member 名称；不支持模糊匹配、目录或递归 ZIP")
            blob = archive.read(entry)
        text, encoding = decode_text(blob)
        result.update(member=member, member_sha256=hashlib.sha256(blob).hexdigest(), encoding=encoding)
    else:
        if member:
            raise ValueError("只有 ZIP 资料包需要 member 参数")
        if meta.get("status") != "ready":
            raise ValueError("附件尚未完成文字提取，请等待或重新上传")
        text = meta.get("text", "")
        result["method"] = meta.get("method", "")
    if action == "list":
        return {**result, "full_chars": len(text), "complete": True,
                "next_step": "使用 action=read 与相同 attachment_id 读取文字；全文计算可用 Python 原件副本。"}
    selected = text[offset:offset+limit]
    return {**result, "text": selected, "offset": offset, "returned_chars": len(selected), "full_chars": len(text),
            "start_line": text.count("\n", 0, min(offset, len(text))) + 1,
            "next_offset": offset + len(selected) if offset + len(selected) < len(text) else None,
            "complete": offset == 0 and len(selected) == len(text),
            "next_step": "按 next_offset 继续；片段不能代表全文核验。数字存在不代表样件或条件匹配。"}
