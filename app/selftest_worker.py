"""Frozen-build smoke checks; synthetic fixtures only, no directory scanning."""
from __future__ import annotations

import io
import importlib
import json
import sys
import tempfile
import zipfile
from pathlib import Path

if __name__ == "__main__":
    from . import attachments, tools
    from .config import DEFAULTS, ROOT, VERSION, save_settings, write_json
    from docx import Document
    from openpyxl import Workbook
    from pptx import Presentation
    from pptx.util import Inches
    from PIL import Image, ImageDraw, ImageFont
    import tkinter

    checks = {}
    if (ROOT / "config/settings.json").exists():
        raise ValueError("Self-test requires a fresh, isolated LOCAL_AI_ROOT; refusing to overwrite existing settings")
    checks["folder_picker_runtime"] = tkinter.TkVersion
    write_json(ROOT / "config/settings.json", DEFAULTS)
    checks["default_launch_and_chat_tools"] = DEFAULTS["launch_mode"] == "app" and {"system_time", "ask_user"} <= set(DEFAULTS["tools"]) and DEFAULTS["max_tool_calls"] == 50
    doc = Document(); doc.add_paragraph("Synthetic Word 2026")
    buf = io.BytesIO(); doc.save(buf)
    doc_meta = attachments.upload("synthetic.docx", buf.getvalue())
    result = attachments.extract(doc_meta["id"])
    assert result["status"] == "ready", result
    checks["word_native"] = "Synthetic Word" in result["text"]
    if "--office" in sys.argv:
        doc_meta["read_mode"] = "office"; attachments.save(doc_meta)
        result = attachments.extract(doc_meta["id"])
        checks["word_com"] = result["status"] == "ready" and "Synthetic Word" in result["text"]
        if not checks["word_com"]:
            checks["word_com_error"] = result.get("error")
    book = Workbook(); book.active["A1"] = "Synthetic Excel 2026"
    buf = io.BytesIO(); book.save(buf)
    meta = attachments.upload("synthetic.xlsx", buf.getvalue()); result = attachments.extract(meta["id"])
    checks["excel_native"] = "Synthetic Excel" in result.get("text", "")
    if "--office" in sys.argv:
        meta["read_mode"] = "office"; attachments.save(meta); result = attachments.extract(meta["id"])
        checks["excel_com"] = result["status"] == "ready" and "Synthetic Excel" in result.get("text", "")
        if not checks["excel_com"]:
            checks["excel_com_error"] = result.get("error")
    ppt = Presentation(); slide = ppt.slides.add_slide(ppt.slide_layouts[6])
    slide.shapes.add_textbox(Inches(1), Inches(1), Inches(4), Inches(1)).text = "Synthetic PowerPoint"
    buf = io.BytesIO(); ppt.save(buf)
    meta = attachments.upload("synthetic.pptx", buf.getvalue()); result = attachments.extract(meta["id"])
    checks["ppt_native"] = "Synthetic PowerPoint" in result.get("text", "")
    if "--office" in sys.argv:
        meta["read_mode"] = "office"; attachments.save(meta); result = attachments.extract(meta["id"])
        checks["ppt_com"] = result["status"] == "ready" and "Synthetic PowerPoint" in result.get("text", "")
        if not checks["ppt_com"]:
            checks["ppt_com_error"] = result.get("error")
    im = Image.new("RGB", (900, 180), "white")
    ImageDraw.Draw(im).text((20, 45), "LOCAL TEST 2026", fill="black", font=ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 50))
    buf = io.BytesIO(); im.save(buf, "PNG")
    meta = attachments.upload("synthetic.png", buf.getvalue()); result = attachments.extract(meta["id"])
    checks["local_ocr"] = "2026" in result.get("text", "")
    checks["guide"] = (Path(__file__).resolve().parents[1] / "docs/软件教程.md").is_file()
    package = io.BytesIO()
    with zipfile.ZipFile(package, "w") as archive:
        for file in (Path(__file__).resolve().parents[1] / "plugins/examples/time_demo").glob("*"):
            if file.is_file():
                archive.write(file, file.name)
    tools.import_plugin(package.getvalue())
    save_settings({"plugins": ["time_demo"]})
    tools.save_tool_configuration("time_demo", {"format": "clock"})
    checks["plugin_worker"] = "time" in tools.execute("time_demo", {}, ["time_demo"])
    # Exercise the documented import/export interface in a fresh tool namespace;
    # no files outside this isolated test root are supplied to the plugin.
    package = io.BytesIO()
    with zipfile.ZipFile(package, "w") as archive:
        for file in (Path(__file__).resolve().parents[1] / "plugins/examples/tool_interface_demo").glob("*"):
            if file.is_file():
                archive.write(file, file.name)
    tools.import_plugin(package.getvalue())
    save_settings({"plugins": ["time_demo", "tool_interface_demo"]})
    tools.save_tool_configuration("tool_interface_demo", {"prefix": "SYNTHETIC-PREFIX:"})
    checks["interface_demo_worker"] = tools.execute("tool_interface_demo", {"action": "calculate", "operation": "multiply", "a": 6, "b": 7}, ["tool_interface_demo"])["result"] == 42
    exported = tools.export_plugin("tool_interface_demo")
    copied = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(exported)) as source, zipfile.ZipFile(copied, "w") as target:
        checks["tool_export_excludes_runtime_settings"] = not any(name.startswith(("config/", "data/", "logs/")) for name in source.namelist()) and b"SYNTHETIC-PREFIX:" not in b"".join(source.read(name) for name in source.namelist())
        for name in source.namelist():
            blob = source.read(name)
            if name == "manifest.json":
                manifest = json.loads(blob)
                manifest["id"] = "tool_interface_roundtrip"
                blob = json.dumps(manifest, ensure_ascii=False).encode("utf-8")
            target.writestr(name, blob)
    tools.import_plugin(copied.getvalue())
    save_settings({"plugins": ["time_demo", "tool_interface_demo", "tool_interface_roundtrip"]})
    checks["tool_export_import_roundtrip"] = tools.execute("tool_interface_roundtrip", {"action": "echo", "text": "SYNTHETIC"}, ["tool_interface_roundtrip"])["text"] == "示例：SYNTHETIC"
    try:
        tools.execute("ask_user", {"question": "SYNTHETIC QUESTION"}, ["ask_user"])
        checks["ask_user_test_does_not_wait"] = False
    except ValueError as error:
        checks["ask_user_test_does_not_wait"] = "聊天" in str(error)
    # Imports and pure helper checks verify frozen dependency collection without
    # starting Office, reading user documents or requiring Office activation.
    office_dependencies = ("pythoncom", "pywintypes", "win32com.client", "win32com.shell.shell",
                           "win32file", "win32con", "win32clipboard", "win32process", "pypdfium2")
    checks["office_com_and_preview_dependencies"] = all(importlib.import_module(name) is not None for name in office_dependencies)
    from . import office_operations, office_worker, snapshots, audit
    checks["office_operator_help"] = "help" in tools.execute("office_operator", {"action": "help", "application": "word"}, ["office_operator"]) and office_worker.color("#FF0000") == 255 and callable(office_operations.execute)
    synthetic_files = ROOT / "data/selftest-files"
    synthetic_files.mkdir()
    create_args = {"directory": str(synthetic_files), "name": "synthetic", "format": "md", "text": "# SYNTHETIC ORIGINAL"}
    created = tools.execute("text_create", create_args, ["text_create"])
    duplicate = tools.execute("text_create", create_args, ["text_create"])
    checks["text_create_never_overwrites"] = created["path"] != duplicate["path"] and Path(created["path"]).read_text(encoding="utf-8") == create_args["text"]
    edited = tools.execute("file_manage", {"action": "edit_text", "path": created["path"], "text": "# SYNTHETIC EDITED"}, ["file_manage"])
    checks["file_edit_and_snapshot"] = Path(created["path"]).read_text(encoding="utf-8") == "# SYNTHETIC EDITED" and any(row["id"] == edited["snapshot_id"] for row in snapshots.listing()["items"])
    restored = snapshots.restore(edited["snapshot_id"])
    checks["snapshot_restore"] = restored["ok"] and Path(created["path"]).read_text(encoding="utf-8") == create_args["text"]
    audit_state = audit.listing()
    checks["bounded_audit"] = audit_state["retention_days"] == 365 and audit_state["max_bytes"] == 1_000_000_000 and audit_state["compression"]["format"] == "gzip-batches" and audit_state["total"] >= 5
    from . import python_lab
    tools.save_tool_configuration("python_run", {"enabled": True, "timeout": 2, "output_kb": 4})
    lab = python_lab.execute({"code": 'import statistics\nprint(statistics.mean([2,4,6]))\nprint(input())', "stdin": "SYNTHETIC-IN\n"})
    checks["python_offline_io"] = lab["status"] == "finished" and "4" in lab["stdout"] and "SYNTHETIC-IN" in lab["stdout"]
    checks["python_timeout"] = "超时" in python_lab.execute({"code": "while True: pass"}).get("note", "")
    from . import context, memory, knowledge
    checks["offline_tokenizer"] = context.tokenizer("qwen35") is not None and context.tokenizer("qwen3-embedding") is not None
    checks["unknown_dat"] = attachments.extract(attachments.upload("synthetic.dat", "合成纯文本 DAT".encode())["id"])["status"] == "ready"
    memory.update("# 合成记忆\n优先使用毫米。")
    checks["fixed_memory"] = "毫米" in memory.injection()
    memory.update("# 合成完整重写\n只保留这条。", ai=True)
    memory.update("追加合成内容", mode="append", ai=True)
    checks["memory_revision_free_full_replace"] = memory.read()["text"] == "# 合成完整重写\n只保留这条。\n追加合成内容"
    snapshot = context.runtime_state({"provider": "local"})
    checks["runtime_snapshot_prompt_match"] = snapshot["local_time"] in context.runtime_context({"provider": "local"}, snapshot)[0]["content"] and not snapshot["tool_called"]
    from .api import admin as api_admin, memory as api_memory, knowledge as api_knowledge, interactions as api_interactions, tool_services as api_tools
    checks["api_adapter_modules"] = all(callable(module.build_router) for module in (api_admin, api_memory, api_knowledge, api_interactions, api_tools))
    kb = knowledge.create("合成冻结测试库")
    checks["knowledge_sqlite"] = kb["config"]["model"] == "qwen3-embedding:0.6b"
    from . import skills, store, turn_queue
    chosen = skills.snapshot(["evidence-summary"])
    checks["builtin_skill_and_snapshot"] = len(chosen) == 1 and chosen[0]["instructions"] in skills.injection({"skill_snapshot": chosen}) and [item["id"] for item in skills.listing()["items"] if item["builtin"]] == ["evidence-summary"]
    skill_blob, skill_name, _ = skills.export("evidence-summary")
    imported_skill = skills.import_file(skill_name, skill_blob)
    checks["skill_export_import_roundtrip"] = imported_skill["instructions"] == chosen[0]["instructions"] and imported_skill["id"] != chosen[0]["id"] and not imported_skill["builtin"]
    archive = io.BytesIO()
    with zipfile.ZipFile(archive, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("synthetic-A.dat", "group=A;axis=Z;stiffness=123.45 N/mm\n")
    raw = attachments.extract(attachments.upload("synthetic.zip", archive.getvalue())["id"])
    scope = {"attachment_ids": [raw["id"]], "python_run_ids": []}
    listing = tools.execute("read_attachment", {"attachment_id": raw["id"], "action": "list"}, ["read_attachment"], scope)
    checks["scoped_zip_attachment"] = listing["total"] == 1 and listing["items"][0]["member"] == "synthetic-A.dat"
    tools.save_tool_configuration("python_run", {"enabled": True, "timeout": 10, "output_kb": 8})
    program = '''from pathlib import Path
import zipfile
from openpyxl import Workbook, load_workbook
from workbench import verify_text, record_calculation, export_evidence
proof=verify_text(INPUTS[0]['id'],'stiffness=123.45 N/mm',member='synthetic-A.dat')
assert proof['occurrences']==1
assert verify_text(INPUTS[0]['id'],'stiffness=999 N/mm',member='synthetic-A.dat')['status']=='not_found'
w=Workbook();w.active.append(['来源','实测N/mm','状态']);w.active.append(['synthetic-A.dat / group A',123.45,'条件待确认']);w.save('synthetic-result.xlsx')
r=load_workbook('synthetic-result.xlsx',read_only=True);assert r.active['B2'].value==123.45;r.close()
export_evidence('synthetic-evidence.json')
print('FROZEN_XLSX_REOPEN_OK')
'''
    generated = python_lab.execute({"code": program}, scope)
    checks["python_attachment_evidence_and_excel_reopen"] = generated["status"] == "finished" and "FROZEN_XLSX_REOPEN_OK" in generated["stdout"]
    checks["python_generated_files_listed"] = {"synthetic-result.xlsx", "synthetic-evidence.json"} <= {p["name"] for p in generated["files"]}
    with store.connect() as db:
        tables = {r[0] for r in db.execute("SELECT name FROM sqlite_master WHERE type='table'")}
    checks["persistent_turn_queue_schema"] = "chat_turns" in tables
    from . import service_config
    checks["port_roles"] = (service_config.GATEWAY_PORT, service_config.CHAT_PORT, service_config.EMBED_PORT) == (11436, 11437, 11438)
    toolkit = Path(__file__).resolve().parent / "toolkit"
    checks["modular_tool_packages"] = all((toolkit / name / "README.md").is_file() for name in ("system_time", "list_directory", "file_info", "search_files", "read_file", "read_attachment", "python_run", "ask_user", "text_create", "file_manage", "http_request", "office_operator"))
    checks["network_tool_available_without_request"] = any(item["function"]["name"] == "http_request" for item in tools.definitions(["http_request"]))
    before_memory = memory.read()["text"]
    try:
        memory.update("SYNTHETIC " * 1200)
        checks["memory_token_hard_limit"] = False
    except memory.MemoryOperationError as error:
        checks["memory_token_hard_limit"] = error.code == "memory_token_limit" and memory.read()["text"] == before_memory
    doc = knowledge.create_markdown(kb["id"], "synthetic.md", "# 合成设备\n首版信息")
    knowledge.update_document_source(kb["id"], doc["id"], {"name": "synthetic-renamed.md", "text": "# 合成设备\n第二版信息"})
    source = knowledge.document_source(kb["id"], doc["id"])
    checks["markdown_source_edit_lifecycle"] = source["name"] == "synthetic-renamed.md" and "第二版" in source["text"]
    knowledge.delete_document(kb["id"], doc["id"])
    output = json.dumps({"version": VERSION, "root": str(ROOT), "frozen": bool(getattr(sys, "frozen", False)), "synthetic_only": True, "checks": checks}, ensure_ascii=False, indent=2)
    (ROOT / "logs/frozen-selftest.json").write_text(output, encoding="utf-8")
    print(output, flush=True)
    if not all(v for k, v in checks.items() if not k.endswith("_error")):
        raise SystemExit(1)
