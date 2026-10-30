"""Synthetic-only regression tests. Never reads existing user documents."""
import asyncio
import io
import json
import os
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

TEST_ROOT = Path(tempfile.mkdtemp(prefix="local-ai-test-"))
os.environ["LOCAL_AI_ROOT"] = str(TEST_ROOT)
SOURCE_ROOT = Path(__file__).resolve().parents[1]
import sys
sys.path.insert(0, str(SOURCE_ROOT))

import httpx
from fastapi.testclient import TestClient
from app import attachments, chat, dialogs, providers, store, tools
from app.config import DEFAULTS, ROOT, save_settings, settings, write_json, secret
from app.server import app


class CoreTests(unittest.TestCase):
    def setUp(self):
        write_json(ROOT / "config/settings.json", DEFAULTS)
        write_json(ROOT / "config/tools.json", {})
        self.client = TestClient(app)

    def test_security(self):
        self.assertEqual(self.client.get("/api/conversations").status_code, 200)
        self.assertEqual(self.client.post("/api/conversations", headers={"Origin": "https://example.invalid"}).status_code, 403)
        self.assertEqual(self.client.get("/health").status_code, 200)

    def test_branch_restore_export(self):
        conv = self.client.post("/api/conversations").json()
        c = store.get(conv["id"])
        first = store.append(c, "user", "Synthetic question")
        old_answer = store.append(c, "assistant", "Original answer")
        store.put(c)
        result = self.client.post(f'/api/conversations/{c["id"]}/message-action', json={"action": "edit", "message_id": first["id"], "content": "Edited question"}).json()
        self.assertEqual(len(store.get(c["id"])["messages"]), 3)
        self.assertEqual(result["message_page"]["total"], 1)
        self.assertEqual(store.path(result)[-1]["content"], "Edited question")
        result = self.client.patch(f'/api/conversations/{c["id"]}', json={"leaf": old_answer["id"]}).json()
        self.assertEqual(store.path(result)[-1]["content"], "Original answer")
        self.assertIn("Synthetic question", self.client.get(f'/api/conversations/{c["id"]}/export').text)

    def test_unrestricted_directory_with_os_path_checks(self):
        allowed = ROOT / "synthetic_files"
        allowed.mkdir(exist_ok=True)
        (allowed / "test.txt").write_text("Synthetic text", encoding="utf-8")
        save_settings({"roots": [{"id": "demo", "path": str(allowed)}]})
        self.assertEqual(tools.execute("list_directory", {"root_id": "demo"}, ["list_directory"])["items"][0]["name"], "test.txt")
        outside = ROOT / "synthetic_outside_alias"
        outside.mkdir(exist_ok=True)
        (outside / "visible.txt").write_text("Synthetic outside alias", encoding="utf-8")
        result = tools.execute("list_directory", {"path": str(outside)}, ["list_directory"])
        self.assertEqual(result["items"][0]["name"], "visible.txt")
        for invalid in ({"path": "relative_only"}, {"root_id": "missing_alias"}):
            with self.assertRaises(ValueError):
                tools.execute("list_directory", invalid, ["list_directory"])
        with self.assertRaises(ValueError):
            tools.execute("system_time", {}, [])

    def test_model_link_validation(self):
        self.assertEqual(providers.parse_model("https://ollama.com/library/qwen3.5:4b"), "qwen3.5:4b")
        for link in ("https://example.invalid/model", "a;whoami", "../../x", "https://ollama.com.evil/library/a"):
            with self.assertRaises(ValueError):
                providers.parse_model(link)

    def test_directory_picker_selection_and_cancel_are_not_authorization(self):
        with patch.object(dialogs, "choose_directory", return_value=str(ROOT)):
            response = self.client.post("/api/ui/pick-directory")
        self.assertEqual(response.json()["path"], str(ROOT))
        self.assertEqual(settings()["roots"], [])
        with patch.object(dialogs, "choose_directory", return_value=""):
            self.assertEqual(self.client.post("/api/ui/pick-directory").json(), {"path": ""})
        with patch.object(dialogs, "choose_directory", side_effect=ValueError("选择窗口已打开")):
            self.assertEqual(self.client.post("/api/ui/pick-directory").status_code, 400)
        with patch.object(dialogs, "choose_directory", return_value=""):
            self.assertEqual(self.client.post("/api/ui/pick-directory").status_code, 200)

    def test_tool_configuration_search_and_read(self):
        base = ROOT / "folder_fixture"
        (base / "one/two").mkdir(parents=True, exist_ok=True)
        for rel in ("sample.txt", "one/sample-child.txt", "one/two/sample-grandchild.txt"):
            (base / rel).write_text("Synthetic text example", encoding="utf-8")
        roots = [{"id": "demo", "path": str(base)}]
        saved = self.client.put("/api/tools/search_files/settings", json={"roots": roots, "search_depth": 1})
        self.assertEqual(saved.status_code, 200, saved.text)
        self.assertEqual(settings()["roots"], roots)
        self.assertEqual(len(tools.search_files({"root_id": "demo", "query": "sample"})["items"]), 2)
        self.assertEqual(tools.tool_configuration("read_file")["values"]["tool_read_mode"], "office")
        self.client.put("/api/tools/read_file/settings", json={"roots": roots, "tool_read_mode": "native"})
        result = tools.execute("read_file", {"root_id": "demo", "relative_path": "sample.txt"}, ["read_file"])
        self.assertEqual(result["text"], "Synthetic text example")
        self.assertEqual(tools.file_info({"root_id": "demo", "relative_path": "sample.txt"})["size"], 22)
        self.assertEqual(self.client.put("/api/tools/search_files/settings", json={"roots": roots, "search_depth": 9}).status_code, 400)
        self.assertEqual(self.client.put("/api/tools/list_directory/settings", json={"roots": roots + roots}).status_code, 400)
        self.assertEqual(tools.tool_configuration("read_file")["values"]["tool_read_mode"], "native")

    def test_pagination_and_guide(self):
        for number in range(25):
            c = store.create(settings()); c["title"] = "PaginationFixture-" + str(number); store.put(c)
        data = self.client.get("/api/conversations?search=PaginationFixture-&page=2&page_size=10").json()
        self.assertEqual(data["total"], 25)
        self.assertEqual(len(data["items"]), 10)
        self.assertNotIn("messages", data["items"][0])
        guide = self.client.get("/api/guide").json()["markdown"]
        self.assertIn("选择文件夹", guide)
        self.assertIn("Office COM", guide)
        self.assertEqual(self.client.get("/api/guide?format=md").status_code, 200)

    def test_small_context_reserves_valid_output(self):
        self.assertEqual(chat.output_budget(2048, 2048), 1024)
        chat.check_budget([{"role": "user", "content": "Hello"}], 2048, chat.output_budget(2048))
        with self.assertRaises(ValueError):
            chat.check_budget([{"role": "user", "content": "中" * 3000}], 2048, 1024)

    def test_visible_speed_removes_markdown_markup(self):
        self.assertEqual(chat.visible_chars("# 标题\n\n**文字** [链接](https://example.invalid)"), len("标题文字链接"))

    def test_missing_files_and_invalid_tool_arguments(self):
        base = ROOT / "missing-file-fixture"; base.mkdir(exist_ok=True)
        tools.save_tool_configuration("list_directory", {"roots": [{"id": "demo", "path": str(base)}]})
        response = self.client.post("/api/tools/test", json={"name": "file_info", "arguments": {"root_id": "demo", "relative_path": "absent.txt"}})
        self.assertEqual(response.status_code, 400)
        response = self.client.post("/api/tools/test", json={"name": "system_time", "arguments": {"unexpected": True}})
        self.assertEqual(response.status_code, 400)

    def test_native_documents(self):
        from docx import Document
        from openpyxl import Workbook
        from pptx import Presentation
        from pptx.util import Inches
        doc = Document(); doc.add_paragraph("Synthetic Word 123"); table = doc.add_table(rows=1, cols=2); table.cell(0, 0).text = "TEST"; table.cell(0, 1).text = "42"
        content = io.BytesIO(); doc.save(content)
        a = attachments.upload("demo.docx", content.getvalue()); result = attachments.extract(a["id"])
        self.assertEqual(result["status"], "ready"); self.assertIn("TEST | 42", result["text"])
        book = Workbook(); book.active["A1"] = "Synthetic Excel"; book.active["B1"] = 12
        content = io.BytesIO(); book.save(content)
        a = attachments.upload("demo.xlsx", content.getvalue()); result = attachments.extract(a["id"])
        self.assertIn("Synthetic Excel | 12", result["text"])
        ppt = Presentation(); slide = ppt.slides.add_slide(ppt.slide_layouts[6]); slide.shapes.add_textbox(Inches(1), Inches(1), Inches(5), Inches(1)).text = "Synthetic Slide"
        content = io.BytesIO(); ppt.save(content)
        a = attachments.upload("demo.pptx", content.getvalue()); result = attachments.extract(a["id"])
        self.assertIn("Synthetic Slide", result["text"])
        a = attachments.upload("demo.txt", "Synthetic UTF8 中文".encode("utf-8")); result = attachments.extract(a["id"])
        self.assertIn("中文", result["text"])
        snap = attachments.snapshot(a["id"], "Selected text", "text")
        self.assertEqual(snap["text"], "Selected text"); self.assertNotEqual(result["text"], snap["text"])

    def test_ocr_synthetic_image(self):
        from PIL import Image, ImageDraw, ImageFont
        im = Image.new("RGB", (900, 180), "white")
        font = ImageFont.truetype("C:/Windows/Fonts/arial.ttf", 45)
        ImageDraw.Draw(im).text((30, 55), "LOCAL TEST 2026", fill="black", font=font)
        content = io.BytesIO(); im.save(content, "PNG")
        a = attachments.upload("synthetic.png", content.getvalue()); result = attachments.extract(a["id"])
        self.assertEqual(result["status"], "ready", result.get("error"))
        self.assertIn("2026", result["text"])
        self.assertTrue((attachments.directory(a["id"]) / "vision.jpg").is_file())

    def test_plugin_import_traversal_and_execution(self):
        bad = io.BytesIO()
        manifest = {"id": "unsafe_demo", "name": "Test", "version": "1", "description": "Synthetic", "parameters": {"type": "object"}, "permissions": []}
        with zipfile.ZipFile(bad, "w") as z:
            z.writestr("manifest.json", json.dumps(manifest)); z.writestr("../tool.py", "pass")
        with self.assertRaises(ValueError):
            tools.import_plugin(bad.getvalue())
        content = io.BytesIO()
        with zipfile.ZipFile(content, "w") as z:
            for file in (SOURCE_ROOT / "plugins/examples/time_demo").glob("*"):
                z.write(file, file.name)
        manifest = tools.import_plugin(content.getvalue())
        with self.assertRaises(ValueError):
            tools.execute(manifest["id"], {}, [manifest["id"]])
        save_settings({"plugins": ["time_demo"]})
        tools.save_tool_configuration("time_demo", {"format": "clock"})
        with patch.dict(os.environ, {"PYTHONPATH": str(SOURCE_ROOT)}):
            result = tools.execute("time_demo", {}, ["time_demo"])
        self.assertIn("time", result)
        self.assertRegex(result["time"], r"^\d{2}:\d{2}:\d{2}$")

    def test_reset_keeps_models_without_backup(self):
        (ROOT / "models/synthetic-sentinel").write_text("Synthetic model marker")
        c = store.create(settings()); store.append(c, "user", "Synthetic reset"); store.put(c)
        response = self.client.post("/api/factory-reset", json={"confirm": "初始化", "clear_chats": True, "clear_providers": True})
        self.assertEqual(response.status_code, 200, response.text)
        self.assertTrue((ROOT / "models/synthetic-sentinel").exists())
        self.assertIsNone(response.json()["backup"])
        self.assertFalse(list((ROOT / "data/backups").iterdir()))
        self.assertEqual(store.listing(), [])
        self.assertEqual(settings()["theme"], "dark")

    def test_key_encryption(self):
        save_settings({"providers": [{"id": "fake", "name": "Synthetic", "base_url": "https://example.invalid/v1", "models": ["test"], "api_key": "SYNTHETIC-NOT-A-REAL-KEY"}]})
        self.assertNotIn("SYNTHETIC-NOT-A-REAL-KEY", (ROOT / "config/secrets.json").read_text())
        self.assertEqual(secret("fake"), "SYNTHETIC-NOT-A-REAL-KEY")

    def test_json_decision_rejects_unexpected_tool(self):
        self.assertEqual(chat.decode_decision('{"action":"answer"}', ["system_time"])["action"], "answer")
        for raw in ('{"action":"call_tool","name":"shell","arguments":{}}', 'not json', '{"action":"execute"}'):
            with self.assertRaises((ValueError, json.JSONDecodeError)):
                chat.decode_decision(raw, ["system_time"])

    def test_native_and_legacy_use_same_executor(self):
        async def exercise(mode):
            c = store.create(settings()); c["options"]["tool_mode"] = mode; store.append(c, "user", "What time is it?")
            count = 0
            async def stream(ident, body):
                nonlocal count
                count += 1
                if mode == "native" and count == 1:
                    yield {"message": {"tool_calls": [{"function": {"name": "system_time", "arguments": {}}}]}, "done": True}
                else:
                    yield {"message": {"content": "Synthetic answer"}, "done": True, "eval_count": 5, "eval_duration": 1000000000}
            decisions = iter(['{"action":"call_tool","name":"system_time","arguments":{}}', '{"action":"answer"}'])
            async def decide(*args):
                return next(decisions)
            with patch.object(providers, "stream_model", stream), patch.object(providers, "decision", decide):
                packets = [p async for p in chat.run(c, {"vision": False, "tools": True}, asyncio.Event())]
            tool_events = [p for p in packets if p["type"] == "tool"]
            self.assertEqual(len(tool_events), 1); self.assertEqual(tool_events[0]["name"], "system_time")
            self.assertTrue(tool_events[0]["success"])
        asyncio.run(exercise("native")); asyncio.run(exercise("legacy"))

    def test_openai_stream_adapter(self):
        async def exercise():
            save_settings({"providers": [{"id": "mock_cloud", "name": "Synthetic", "base_url": "https://example.invalid/v1", "models": ["test"], "tools": True}]})
            def handler(request):
                payload = json.loads(request.content)
                self.assertEqual(payload["model"], "test")
                self.assertEqual(request.url.path, "/v1/chat/completions")
                body = 'data: {"choices":[{"delta":{"content":"Hi"}}]}\n\ndata: {"choices":[],"usage":{"prompt_tokens":3,"completion_tokens":1}}\n\ndata: [DONE]\n\n'
                return httpx.Response(200, text=body, headers={"content-type": "text/event-stream"})
            def client(*args):
                return httpx.AsyncClient(base_url="https://example.invalid/v1/", transport=httpx.MockTransport(handler))
            with patch.object(providers, "cloud_client", client):
                packets = [p async for p in providers.stream_model("mock_cloud", {"model": "test", "messages": [{"role": "user", "content": "Test"}]})]
            self.assertEqual(packets[0]["message"]["content"], "Hi")
            self.assertEqual(packets[1]["usage"]["completion_tokens"], 1)
        asyncio.run(exercise())

    def test_ollama_cloud_alias_cannot_masquerade_as_local(self):
        async def remote(*args, **kwargs):
            return {"remote_model": "synthetic-cloud-model", "remote_host": "https://example.invalid", "capabilities": ["completion"]}
        with patch.object(providers, "local", remote):
            with self.assertRaises(ValueError):
                asyncio.run(providers.capability("local", "synthetic-cloud-alias"))


if __name__ == "__main__":
    unittest.main(verbosity=2)
