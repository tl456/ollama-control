"""Synthetic-only skill, attachment scope, and trusted Python provenance regressions."""
import asyncio
import hashlib
import io
import json
import stat
import unittest
import zipfile
from unittest.mock import patch

from test_core import ROOT
from app import attachments, chat, context, providers, python_lab, skills, store, tools
from app.config import DEFAULTS, settings, write_json


class SkillTests(unittest.TestCase):
    def setUp(self):
        write_json(ROOT / "config/settings.json", DEFAULTS)
        write_json(ROOT / "config/tools.json", {})

    def text(self, name="synthetic.dat", text="组=A;方向=X;频率=10Hz;振幅=±0.1mm;刚度=101N/mm\n组=B;方向=Y;频率=20Hz;刚度=101N/mm\n"):
        value = attachments.upload(name, text.encode("utf-8"))
        return attachments.extract(value["id"])

    def archive(self, entries):
        blob = io.BytesIO()
        with zipfile.ZipFile(blob, "w", zipfile.ZIP_DEFLATED) as archive:
            for name, value in entries:
                archive.writestr(name, value)
        return blob.getvalue()

    def scope(self, meta):
        return {"attachment_ids": [meta["id"]], "python_run_ids": []}

    def test_skill_discovery_and_short_on_demand_rules(self):
        listing = skills.listing()
        self.assertEqual(listing["max_selected"], 2)
        item = next(v for v in listing["items"] if v["id"] == "evidence-summary")
        self.assertNotIn("instructions", item)
        self.assertEqual(skills.injection({}), "")
        loaded = skills.injection({"skill_ids": [item["id"]]})
        self.assertIn("不要编造缺失数据", loaded)
        self.assertIn("待确认", loaded)
        self.assertLess(context.estimate(loaded), 1800)
        self.assertEqual(item["requires_tools"], ["read_attachment", "python_run"])

    def test_skill_snapshot_immutable_and_integrity_checked(self):
        snapshot = skills.snapshot(["evidence-summary"])
        with patch.object(skills, "_read", side_effect=ValueError("changed")):
            self.assertIn("自主", skills.injection({"skill_snapshot": snapshot}))
        snapshot[0]["instructions"] += "变更"
        with self.assertRaisesRegex(ValueError, "快照"):
            skills.injection({"skill_snapshot": snapshot})
        for ids in (["evidence-summary"] * 2, ["missing"], "evidence-summary", ["../secret"]):
            with self.assertRaises(ValueError):
                skills.snapshot(ids)

    def test_context_only_selected_skill_and_attachment_ids(self):
        meta = self.text()
        conv = store.create(settings())
        store.append(conv, "user", "整理", attachments=[attachments.snapshot(meta["id"])])
        before = context.history(conv, False)
        self.assertFalse(any("已选择技能" in m["content"] for m in before))
        conv["options"]["skill_ids"] = ["evidence-summary"]
        after = context.history(conv, False)
        self.assertTrue(any("已选择技能" in m["content"] for m in after))
        self.assertIn(meta["id"], after[-1]["content"])
        self.assertNotIn("python_run", conv["options"]["tools"])

    def test_read_attachment_requires_branch_scope(self):
        current, other = self.text(), self.text("elsewhere.dat")
        conv = store.create(settings())
        first = store.append(conv, "user", "本轮", attachments=[attachments.snapshot(current["id"])])
        store.append(conv, "user", "未来分支", attachments=[attachments.snapshot(other["id"])])
        conv["leaf"] = first["id"]
        scope = chat.tool_scope(conv)
        self.assertEqual(scope["attachment_ids"], [current["id"]])
        result = tools.execute("read_attachment", {"action": "list"}, ["read_attachment"], scope)
        self.assertEqual(result["items"][0]["id"], current["id"])
        with self.assertRaisesRegex(ValueError, "当前对话"):
            tools.execute("read_attachment", {"action": "read", "attachment_id": other["id"]}, ["read_attachment"], scope)
        first["excluded"] = True
        self.assertEqual(chat.tool_scope(conv)["attachment_ids"], [])
        self.assertEqual(tools.execute("read_attachment", {}, ["read_attachment"])["items"], [])

    def test_attachment_segment_resume_and_source_integrity(self):
        original = "字段\n" + "ABCDE" * 80
        meta = self.text(text=original)
        args = {"action": "read", "attachment_id": meta["id"], "max_chars": 100}
        first = tools.execute("read_attachment", args, ["read_attachment"], self.scope(meta))
        rest = tools.execute("read_attachment", {**args, "offset": first["next_offset"], "max_chars": 12000}, ["read_attachment"], self.scope(meta))
        self.assertFalse(first["complete"])
        self.assertEqual(first["text"] + rest["text"], original)
        self.assertIsNone(rest["next_offset"])
        _, source = attachments.raw_source(meta["id"])
        source.write_text("已修改", encoding="utf-8")
        with self.assertRaisesRegex(ValueError, "校验变化"):
            tools.execute("read_attachment", args, ["read_attachment"], self.scope(meta))

    def test_zip_raw_attachment_list_and_read(self):
        raw = self.archive([("series/part-A.dat", "组A,方向X,101\n"), ("series/part-B.dat", "组B,方向Y,101\n")])
        meta = attachments.upload("synthetic.zip", raw)
        meta = attachments.extract(meta["id"])
        self.assertEqual(meta["status"], "ready")
        self.assertTrue(meta["raw_archive"])
        self.assertNotIn("组A,方向X", meta["text"])
        snap = attachments.snapshot(meta["id"])
        self.assertTrue(snap["raw_archive"])
        result = attachments.read_authorized({"action": "list", "attachment_id": meta["id"]}, self.scope(meta))
        self.assertEqual(result["total"], 2)
        text = attachments.read_authorized({"action": "read", "attachment_id": meta["id"], "member": result["items"][0]["member"]}, self.scope(meta))
        self.assertEqual(text["text"], "组A,方向X,101\n")
        self.assertTrue(text["complete"])
        with self.assertRaises(ValueError):
            attachments.read_authorized({"action": "read", "attachment_id": meta["id"], "member": "../part-A.dat"}, self.scope(meta))
        self.assertFalse((attachments.directory(meta["id"]) / "series").exists())

    def test_zip_entry_pagination(self):
        raw = self.archive([(f"record-{n}.dat", f"record {n}") for n in range(101)])
        meta = attachments.extract(attachments.upload("many.zip", raw)["id"])
        args = {"action": "list", "attachment_id": meta["id"]}
        first = attachments.read_authorized(args, self.scope(meta))
        last = attachments.read_authorized({**args, "offset": first["next_offset"]}, self.scope(meta))
        self.assertEqual(len(first["items"]), 100)
        self.assertEqual(len(last["items"]), 1)
        self.assertIsNone(last["next_offset"])

    def test_zip_unsafe_paths_duplicates_and_ratio_rejected(self):
        for name in ("../outside.dat", "/absolute.dat", "C:/drive.dat", "dir/./item.dat"):
            with self.subTest(name=name), self.assertRaises(ValueError):
                attachments.upload("bad.zip", self.archive([(name, "x")]))
        # ZipInfo normalizes os.sep on Windows, so replace equal-length names in both headers.
        raw_backslash = self.archive([("dir/item.dat", "x")]).replace(b"dir/item.dat", b"dir\\item.dat")
        with self.assertRaises(ValueError):
            attachments.upload("bad.zip", raw_backslash)
        with self.assertRaisesRegex(ValueError, "重复"):
            attachments.upload("dup.zip", self.archive([("A.dat", "one"), ("a.dat", "two")]))
        with self.assertRaisesRegex(ValueError, "压缩比"):
            attachments.upload("ratio.zip", self.archive([("zero.dat", "0" * (2 * 1024 * 1024))]))
        blob = io.BytesIO()
        with zipfile.ZipFile(blob, "w") as archive:
            item = zipfile.ZipInfo("link.dat")
            item.external_attr = (stat.S_IFLNK | 0o777) << 16
            archive.writestr(item, "target")
        with self.assertRaisesRegex(ValueError, "符号链接"):
            attachments.upload("link.zip", blob.getvalue())

    def test_python_raw_inputs_evidence_and_excel_roundtrip(self):
        meta = self.text()
        code = '''from pathlib import Path
import json
from openpyxl import Workbook, load_workbook
from workbench import verify_text, record_calculation, export_evidence
print(INPUTS[0]['id'])
raw = Path(INPUTS[0]['path']).read_text(encoding='utf-8')
assert '组=B' in raw
check = verify_text(INPUTS[0]['id'], '组=A;方向=X;频率=10Hz;振幅=±0.1mm;刚度=101N/mm')
missing = verify_text(INPUTS[0]['id'], '刚度=999N/mm')
assert check['occurrences'] == 1 and missing['status'] == 'not_found'
assert '未自动核验' in check['semantic_status']
record_calculation('相对目标偏差', {'measured':101,'target':100}, '(measured-target)/target', .01, [check])
export_evidence()
book = Workbook(); sheet = book.active; sheet.title = '结果'
sheet.append(['组', '刚度N/mm', '条件状态']); sheet.append(['A',101,'标准待确认'])
book.save('合成结果.xlsx')
opened = load_workbook('合成结果.xlsx', read_only=True)
assert opened['结果']['B2'].value == 101; opened.close()
print('ROUNDTRIP_OK')'''
        result = tools.execute("python_run", {"code": code}, ["python_run"], self.scope(meta))
        self.assertEqual(result["exit_code"], 0, result["stderr"])
        self.assertIn("ROUNDTRIP_OK", result["stdout"])
        files = {v["name"]: v for v in result["files"]}
        self.assertIn("合成结果.xlsx", files)
        self.assertIn("复核记录.json", files)
        self.assertNotIn("workbench.py", files)
        evidence = json.loads((python_lab.folder(result["id"]) / "复核记录.json").read_text(encoding="utf-8"))
        self.assertEqual(evidence["records"][0]["status"], "literal_found")
        self.assertEqual(evidence["records"][1]["status"], "not_found")
        self.assertIn("非独立重算", evidence["records"][2]["status"])
        with (python_lab.folder(result["id"]) / "合成结果.xlsx").open("rb") as reader:
            self.assertEqual(hashlib.file_digest(reader, "sha256").hexdigest(), files["合成结果.xlsx"]["sha256"])
        continuation = tools.execute("python_run", {"previous_run_id": result["id"], "code": "from openpyxl import load_workbook\nb=load_workbook('previous/合成结果.xlsx');print(b.active['B2'].value);b.close()"},
                                     ["python_run"], {**self.scope(meta), "python_run_ids": [result["id"]]})
        self.assertEqual(continuation["exit_code"], 0, continuation["stderr"])
        self.assertEqual(continuation["stdout"].strip(), "101")

    def test_python_zip_original_and_missing_scope_denied(self):
        meta = attachments.extract(attachments.upload("raw.zip", self.archive([("data.dat", "组=A;结果=101")]))["id"])
        code = "from workbench import verify_text\nprint(verify_text(INPUTS[0]['id'],'结果=101',member='data.dat')['status'])"
        result = python_lab.execute({"code": code}, self.scope(meta))
        self.assertEqual(result["exit_code"], 0, result["stderr"])
        self.assertIn("literal_found", result["stdout"])
        for args in ({"code": "print(1)", "attachment_ids": [meta["id"]]},
                     {"code": "print(1)", "previous_run_id": result["id"]}):
            with self.assertRaisesRegex(ValueError, "当前对话"):
                python_lab.execute(args, {})

    def test_tool_errors_are_short_and_do_not_echo_code(self):
        async def run():
            return await chat.call_tool("python_run", {"code": "SECRET_BODY", "unknown": "private"}, ["python_run"], "native")
        error = asyncio.run(run())
        self.assertFalse(error["success"])
        self.assertNotIn("SECRET_BODY", error["result"]["error"])
        self.assertIn("参数", error["result"]["error"])
        self.assertIn("next_step", error["result"])
        self.assertFalse(error["result"]["executed"])

    def test_archive_resource_limits_without_expanding_payloads(self):
        class Directory:
            def __init__(self, values):
                self.values = values
            def infolist(self):
                return self.values
        item = zipfile.ZipInfo("data.dat")
        item.file_size = 21 * 1024 * 1024
        item.compress_size = item.file_size
        with self.assertRaisesRegex(ValueError, "20 MiB"):
            attachments.archive_entries(Directory([item]))
        item.file_size = 1
        with self.assertRaisesRegex(ValueError, "1 万"):
            attachments.archive_entries(Directory([item] * 10001))
        item.file_size = 101 * 1024 * 1024
        with self.assertRaisesRegex(ValueError, "100 MiB"):
            attachments.archive_entries(Directory([item]))

    def test_native_skill_rounds_bounded_and_no_implicit_python(self):
        write_json(ROOT / "config/settings.json", {**DEFAULTS, "memory_enabled": False, "context": 65536})
        async def fake_stream(provider, body, stop):
            if body.get("tools"):
                yield {"message": {"tool_calls": [{"function": {"name": "system_time", "arguments": {}}}]}, "done": True}
            else:
                yield {"message": {"content": "来源仍需复核"}, "done": True}
        async def run(selected):
            conv = store.create(settings())
            conv["options"].update(tools=["system_time"], skill_ids=selected)
            store.append(conv, "user", "整理并核对")
            with patch.object(providers, "resilient_stream", fake_stream):
                events = [v async for v in chat.run(conv, {"tools": True, "vision": False, "max_context": 262144}, asyncio.Event())]
            self.assertFalse(any(v["type"] == "error" for v in events), events)
            self.assertNotIn("python_run", conv["options"]["tools"])
            return events
        ordinary = asyncio.run(run([]))
        skilled = asyncio.run(run(["evidence-summary"]))
        self.assertEqual(sum(v["type"] == "tool" for v in ordinary), 50)
        self.assertEqual(sum(v["type"] == "tool" for v in skilled), 50)
        self.assertTrue(any(v["type"] == "phase" and "已装载技能" in v["text"] for v in skilled))


if __name__ == "__main__":
    unittest.main()
